from datetime import date
import json
import zipfile

import openpyxl
import pytest

from excel_helper.fill_workbook import address, richgo_facts, PublicPages, fill
from estate_ui_release import copy_ui


def test_asking_range_and_recency_are_not_transaction_or_predicted_prices():
    data = {'basic': {'construction': {'date': '1994.04.21'}, 'households': 498},
            'pyeongInfos': {'30': {'pyeongType': 30, 'minPrivateArea': '84.07', 'maxPrivateArea': '84.34',
                'danjiPriceInfo': {'memePriceDict': {'MOLIT': {'price': 74000},
                    'RICHGO_SISE': {'price': 120000},
                    'OFFER': {'minPrice': 133000, 'yyyymmdd': '2026.10.07'}}}}}}
    got = richgo_facts(data, 84.34, date(2026, 10, 8))
    assert got['매물 최저호가(억)'] == 13.3
    assert '84.07~84.34' in got['호가 평형 범위']
    assert '매물 최저호가(억)' not in richgo_facts(data, 84.77, date(2026, 10, 8))
    assert '매물 최저호가(억)' not in richgo_facts(data, 84.34, date(2026, 11, 8))
    assert address('서울특별시 강서구 염창동 288') == address('서울시 강서구 염창동 288')
    assert address('서울 강서구 염창동 288') != address('서울 강서구 염창동 291')


def test_autofill_preserves_manual_cells_original_and_source_trace(tmp_path, monkeypatch):
    import excel_helper.fill_workbook as helper
    src = tmp_path/'input.xlsx'
    w = openpyxl.Workbook(); s = w.active; s.title = '임장 비교'
    s.append(['단지', '주소', '전용(㎡)', '준공년도', '용적률(%)', '임장', '메모'])
    s.append(['현대1', '서울 강서구 염창동 288', 84.34, 1994, None, '방문함', '=이 문자는 메모'])
    w.create_sheet('기준 및 출처').append(['항목', '설명'])
    w.save(src); original = src.read_bytes()
    monkeypatch.setattr(helper, 'collect', lambda *a: ({'용적률(%)': 271, '임장': '덮어쓰기 금지'}, ['https://m.richgo.ai/realty/danji/example'], []))
    result = fill(src)
    out = openpyxl.load_workbook(result['output'])
    assert src.read_bytes() == original
    assert out.worksheets[0]['E2'].value == 271
    assert out.worksheets[0]['F2'].value == '방문함'
    assert out.worksheets[0]['G2'].value == '=이 문자는 메모'
    assert 'https://m.richgo.ai/' in out.worksheets[0]['E2'].comment.text
    assert out.sheetnames == ['임장 비교', '기준 및 출처']


def test_helper_zip_contains_only_program_and_instructions(tmp_path):
    copy_ui(tmp_path)
    with zipfile.ZipFile(tmp_path/'downloads/estate-excel-helper.zip') as z:
        assert {'fill.cmd', 'fill_workbook.py', 'requirements.txt', '사용법.md', 'FillEstate.bas'} <= set(z.namelist())
        assert all('/' not in n and not n.endswith(('.xlsx', '.html', '.json')) for n in z.namelist())


def test_naver_time_includes_wait_and_rejects_missing_destination():
    from excel_helper.naver_routes import arrival_minutes, destination_token
    assert arrival_minutes('최적 오전 8:05 출발 오전 8:27 도착') == 27
    assert arrival_minutes('오후 1:15 도착') == 315
    with pytest.raises(ValueError):
        arrival_minutes('오전 7:50 도착')
    with pytest.raises(ValueError):
        destination_token('https://map.naver.com/p/directions/-/-/-/transit')
    with pytest.raises(ValueError):
        destination_token('https://example.com/p/directions/-/destination/-/transit')
