"""임장 비교표 자동 채우기. 인증키 없이 공개 페이지를 조회하며 원본은 보존합니다."""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta
import hashlib
import json
from pathlib import Path
import re
import time
from urllib.parse import urlparse, urlencode

from bs4 import BeautifulSoup
import openpyxl
from openpyxl.comments import Comment
import requests

ALLOWED = {'search.naver.com', 'm.richgo.ai', 'realty.daangn.com'}
AUTO = '임장 자동 채우기'


def number(value):
    try:
        v = float(value)
        return v if 0 < v < 1e9 else None
    except (TypeError, ValueError):
        return None


def address(value):
    s = str(value or '').replace('서울특별시', '서울').replace('서울시', '서울').replace('인천광역시', '인천').replace('경기도', '경기')
    return re.sub(r'\s+', ' ', s).strip()


class PublicPages:
    def __init__(self, cache, refresh=False):
        self.cache = Path(cache)
        self.cache.mkdir(parents=True, exist_ok=True)
        self.refresh = refresh
        self.session = requests.Session()
        self.last = 0.

    def get(self, url):
        if urlparse(url).scheme != 'https' or urlparse(url).hostname not in ALLOWED:
            raise ValueError('지원하는 공개 정보 사이트가 아닙니다.')
        cache = self.cache/(hashlib.sha256(url.encode()).hexdigest()+'.html')
        if not self.refresh and cache.exists() and time.time()-cache.stat().st_mtime < 86400:
            return cache.read_text(encoding='utf-8')
        time.sleep(max(0., 1.2-(time.monotonic()-self.last)))
        response = self.session.get(url, timeout=30, allow_redirects=False)
        self.last = time.monotonic()
        response.raise_for_status()
        if response.is_redirect:
            raise ValueError('로그인 또는 변경된 주소로 이동되어 조회를 중단했습니다.')
        if len(response.content) > 8_000_000:
            raise ValueError('페이지 크기 제한을 초과했습니다.')
        body = response.content.decode('utf-8')
        if 'captcha' in body[:2000].lower():
            raise ValueError('사이트의 확인 절차가 필요합니다. 우회하지 않고 건너뜁니다.')
        cache.write_text(body, encoding='utf-8')
        return body

    def search(self, query, provider):
        sites = {'richgo': 'm.richgo.ai/realty/danji', 'daangn': 'realty.daangn.com/complexes'}
        prefix = 'https://'+sites[provider]+'/'
        html = self.get('https://search.naver.com/search.naver?'+urlencode({'query': query+' site:'+sites[provider]}))
        urls = []
        for a in BeautifulSoup(html, 'html.parser').select('a[href]'):
            href = a['href']
            if href.startswith(prefix):
                match = re.match(re.escape(prefix)+r'[a-zA-Z0-9]+', href)
                if match and match[0] not in urls:
                    urls.append(match[0])
        return urls[:6]


def richgo(html):
    node = BeautifulSoup(html, 'html.parser').find('script', id='__NEXT_DATA__')
    if not node:
        raise ValueError('단지 자료를 읽을 수 없습니다.')
    page = json.loads(node.string)['props']['pageProps']
    data = page.get('data') or {}
    if not data.get('basic', {}).get('address', {}).get('jibun'):
        raise ValueError('주소가 있는 단지 자료가 아닙니다.')
    return data


def richgo_facts(data, area, today=None):
    today = today or datetime.now().date()
    basic = data['basic']
    f = {'세대수': number(basic.get('households')), '준공년도': int(basic['construction']['date'][:4]),
         '세대당 주차': round(number(basic.get('parking', {}).get('households')), 3) if number(basic.get('parking', {}).get('households')) else None,
         '용적률(%)': number(basic.get('floorArea')), '건폐율(%)': number(basic.get('buildingRatio'))}
    # 공급평형을 전용면적의 단일 타입 호가로 오인하지 않도록 범위를 함께 기록한다.
    matches = [v for v in data.get('pyeongInfos', {}).values()
               if number(v.get('minPrivateArea')) and number(v.get('maxPrivateArea'))
               and float(v['minPrivateArea'])-.005 <= area <= float(v['maxPrivateArea'])+.005]
    if len(matches) == 1:
        p = matches[0]
        offer = p.get('danjiPriceInfo', {}).get('memePriceDict', {}).get('OFFER', {})
        date_text = offer.get('yyyymmdd') or ''
        try:
            dated = datetime.strptime(date_text, '%Y.%m.%d').date()
        except ValueError:
            dated = None
        if dated and 0 <= (today-dated).days <= 30 and number(offer.get('minPrice')):
            f['매물 최저호가(억)'] = float(offer['minPrice'])/10000
            f['호가 평형 범위'] = f"공급 {p['pyeongType']}평 · 전용 {p['minPrivateArea']}~{p['maxPrivateArea']}㎡"
            f['_asking_date'] = date_text
    return {k: v for k, v in f.items() if v is not None}


def daangn(html, road_address):
    soup = BeautifulSoup(html, 'html.parser')
    script = next((s.string for s in soup.find_all('script') if (s.string or '').startswith('window.RELAY_STORE = ')), None)
    if not script:
        return {}
    # Decode public server-rendered page data as JSON; never execute site scripts.
    encoded, _ = json.JSONDecoder().raw_decode(script.split(' = ', 1)[1])
    data = json.loads(encoded)
    complex_ = next((v for v in data.values() if v.get('__typename') == 'PropComplex' and 'proptierDanji' in v), None)
    if not complex_:
        return {}
    def ref(obj, name):
        return data.get((obj.get(name) or {}).get('__ref'), {})
    def refs(obj, name):
        return [data[k] for k in (obj.get(name) or {}).get('__refs', [])]
    buildings = ref(complex_, 'buildings(first:10)') or ref(complex_, 'buildings(first:1)')
    # Only buildings attached to this complex, never unrelated recommendations.
    linked = [ref(data[e], 'node') for e in (buildings.get('edges') or {}).get('__refs', [])]
    if not any(address(b.get('roadAddress')) == address(road_address) for b in linked):
        return {}
    p = ref(complex_, 'proptierDanji')
    f = {}
    connection = complex_.get('parkingConnection')
    if isinstance(connection, bool):
        f['승강기 연결'] = 'O' if connection else 'X'
        if connection:
            f['지하주차장'] = 'O'
    stations = [s for s in refs(p, 'subways') if number(s.get('walkingTimeMin'))]
    if stations:
        nearest = min(stations, key=lambda s: float(s['walkingTimeMin']))
        f.update({'근접역': nearest['stationNm']+'역', '역 도보(분)': float(nearest['walkingTimeMin'])})
    communities = complex_.get('communityFacilities')
    if isinstance(communities, list) and communities and all(isinstance(c, str) for c in communities):
        labels = {'DAYCARE': '어린이집', 'SENIOR_CENTER': '경로당', 'GYM': '헬스장',
                  'FITNESS': '헬스장', 'GOLF': '골프연습장', 'LIBRARY': '도서관',
                  'SWIMMING_POOL': '수영장', 'SAUNA': '사우나', 'GUEST_HOUSE': '게스트하우스',
                  'PLAYGROUND': '놀이터'}
        translated = [labels.get(c, c if re.search('[가-힣]', c) else '') for c in communities]
        if any(translated):
            f['커뮤니티'] = ', '.join(c for c in translated if c)
    return f


def collect(pages, name, location, area, built):
    locality = location.rsplit(' ', 1)[0]
    short_name = name.removesuffix('아파트')
    dong = locality.split()[-1].removesuffix('동')
    if dong and short_name.startswith(dong) and len(short_name) > len(dong):
        short_name = short_name[len(dong):]
    query = locality+' '+short_name
    errors = []
    for url in pages.search(query, 'richgo'):
        try:
            data = richgo(pages.get(url))
            basic = data['basic']
            if address(basic['address']['jibun']) != address(location):
                continue
            if built and int(basic['construction']['date'][:4]) != int(built):
                continue
            facts = richgo_facts(data, area)
            sources = [url]
            # Confirm the secondary source against the verified road address.
            for other in pages.search(query, 'daangn')[:3]:
                try:
                    additional = daangn(pages.get(other), basic['address']['road'])
                    if additional:
                        facts.update(additional)
                        sources.append(other)
                        break
                except (ValueError, KeyError, requests.RequestException) as e:
                    errors.append(type(e).__name__)
            return facts, sources, errors
        except (ValueError, KeyError, requests.RequestException) as e:
            errors.append(type(e).__name__)
    return {}, [], errors+['주소·준공년도가 일치하는 공개 단지를 찾지 못했습니다.']


def fill(path, output=None, cache=None, refresh=False, limit=None, routes=True):
    path = Path(path).resolve()
    output = Path(output).resolve() if output else path.with_name(path.stem+'-채움'+path.suffix)
    if output == path or output.exists():
        raise ValueError('원본이나 기존 결과를 덮어쓰지 않습니다. 새 출력 이름을 지정하세요.')
    wb = openpyxl.load_workbook(path, keep_vba=path.suffix.lower()=='.xlsm')
    if wb.sheetnames != ['임장 비교', '기준 및 출처']:
        raise ValueError('대시보드에서 최신 임장 비교 엑셀을 내려받아 주세요.')
    ws, notes = wb.worksheets
    headers = {c.value: c.column for c in ws[1]}
    for required in ['단지', '주소', '전용(㎡)', '임장', '메모']:
        if required not in headers:
            raise ValueError('비교표의 필수 열이 없습니다: '+required)
    pages = PublicPages(cache or path.parent/'.임장조회캐시', refresh)
    templates = {str(r[0].value).removeprefix('네이버 목적지 · '): r[1].value
                 for r in notes.iter_rows(min_row=2, max_col=2)
                 if str(r[0].value).startswith('네이버 목적지 · ')}
    route_reader = None
    if routes and templates:
        try:
            from .naver_routes import NaverRoutes
        except ImportError:
            from naver_routes import NaverRoutes
        route_reader = NaverRoutes(pages.cache/'이동시간')
    today = datetime.now().astimezone().isoformat(timespec='seconds')
    changed = processed = 0
    memo = {}
    for row in range(2, ws.max_row+1):
        if limit is not None and processed >= limit:
            break
        values = {h: ws.cell(row, col).value for h, col in headers.items()}
        identity = (values['단지'], values['주소'], number(values['전용(㎡)']), values.get('준공년도'))
        if not identity[0] or not identity[1] or not identity[2]:
            continue
        print(f"[{row-1}/{ws.max_row-1}] {identity[0]} {identity[2]}㎡ 조회", flush=True)
        try:
            if identity not in memo:
                memo[identity] = collect(pages, *identity)
            facts, sources, errors = memo[identity]
        except (requests.RequestException, ValueError, KeyError) as e:
            facts, sources, errors = {}, [], [str(e)[:200]]
        facts, sources, errors = dict(facts), list(sources), list(errors)
        route_notes = []
        if route_reader:
            for label, url in templates.items():
                try:
                    record = route_reader.read(identity[1], url)
                    facts[label+'(분)'] = record['minutes']
                    route_notes.append(f"{label}: {record['route_date']} 08:00 출발 · {record['minutes']}분 · {record['source_url']}")
                except Exception as e:
                    # Site/UI failures cannot invent a time or destroy a workbook.
                    errors.append(label+' 이동 조회 실패: '+str(e).splitlines()[0][:160])
        provenance = f"{today}\n"+'\n'.join(sources)
        if route_notes:
            provenance += '\n'+'\n'.join(route_notes)
        if facts.get('_asking_date'):
            provenance += '\n호가 자료일: '+facts['_asking_date']
        for h, value in facts.items():
            if h.startswith('_') or h not in headers or h in {'임장', '메모'}:
                continue
            cell = ws.cell(row, headers[h])
            # Manual changes to a previously populated result are respected.
            if cell.comment and cell.comment.author == AUTO:
                old = cell.comment.text.split('\n', 1)[0].removeprefix('저장값: ')
                if str(cell.value) != old:
                    continue
            # Existing downloaded dated snapshot values may be refreshed. Values
            # entered in the manual fields are outside this write allowlist.
            cell.value = value
            if isinstance(value, str) and len(value) > 16:
                cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='center')
                ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 30, 48)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.comment = Comment('저장값: '+str(value)+'\n'+provenance, AUTO)
            changed += 1
        if sources and '단지 정보' in headers:
            cell = ws.cell(row, headers['단지 정보'])
            cell.value, cell.hyperlink = '정보 보기', sources[0]
        note = provenance
        if errors:
            note += '\n조회 메모: '+' / '.join(dict.fromkeys(errors))
        note += '\n대지지분·정비사업·미제공 시설은 공개 페이지에서 확정된 값이 없어 보완하지 않았습니다.'
        notes.append([f'자동 조회 {row-1}. {identity[0]} {identity[2]}㎡', note])
        notes.cell(notes.max_row, 2).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        notes.row_dimensions[notes.max_row].height = min(240, 18*(note.count('\n')+2))
        processed += 1
        # A bounded progress file allows recovery without corrupting the source.
        tmp = output.with_name(output.stem+'.진행중'+output.suffix)
        if processed % 10 == 0:
            wb.save(tmp)
    if route_reader:
        route_reader.close()
    wb.save(output)
    progress = output.with_name(output.stem+'.진행중'+output.suffix)
    if progress.exists():
        progress.unlink()
    print(f'완료: {processed}행, {changed}개 항목 → {output}', flush=True)
    return {'rows': processed, 'cells': changed, 'output': str(output)}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('workbook', nargs='?')
    p.add_argument('--output')
    p.add_argument('--cache')
    p.add_argument('--refresh', action='store_true')
    p.add_argument('--limit', type=int)
    p.add_argument('--no-routes', action='store_true', help='시설·호가만 조회')
    a = p.parse_args()
    if not a.workbook:
        import tkinter as tk
        from tkinter.filedialog import askopenfilename
        root = tk.Tk(); root.withdraw()
        a.workbook = askopenfilename(title='채울 임장 비교표 선택', filetypes=[('엑셀', '*.xlsx *.xlsm')])
        root.destroy()
        if not a.workbook:
            return
    fill(a.workbook, a.output, a.cache, a.refresh, a.limit, not a.no_routes)


if __name__ == '__main__':
    main()
